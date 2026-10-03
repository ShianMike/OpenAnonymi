"""Project-authored fictional Word inputs; hidden sentinels prove export isolation."""

import io
import struct
import zlib
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SENTINEL = "ORIGINAL_PRIVATE_SENTINEL"


def numbered_instance(document):
    numbering = document.part.numbering_part.element
    node = numbering.add_num(7)
    node.add_lvlOverride(0).add_startOverride(1)
    return node.numId


def paragraph_number(paragraph, number, level=0):
    properties = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    properties.get_or_add_numId().val = number
    properties.get_or_add_ilvl().val = level


def structured_docx():
    document = Document()
    document.core_properties.author = document.core_properties.title = SENTINEL
    document.sections[0].header.paragraphs[0].text = "Header: nora@example.test"
    document.add_heading("Fictional research", level=0)
    document.add_heading("People and dates", level=1)
    document.add_heading("Deep heading", level=9)
    person = document.add_paragraph("😀 Nora Caldwell met Nora Caldwell.")
    document.add_comment(person.runs[0], SENTINEL, author=SENTINEL, initials="XX")
    document.add_paragraph("First bullet", style="List Bullet")
    document.add_paragraph("Nested bullet", style="List Bullet 3")
    first, second = numbered_instance(document), numbered_instance(document)
    for value in ("First number", "Second number"):
        paragraph_number(document.add_paragraph(value, style="List Number"), first)
    document.add_paragraph("A break between numbered lists.")
    paragraph_number(document.add_paragraph("Restarted number", style="List Number"), second)
    paragraph_number(document.add_paragraph("Level nine number", style="List Number 3"), second, 8)
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Combined cell"
    table.cell(0, 2).text = "nora@example.test"
    table.cell(1, 0).text = "DOB: 2000-01-01"
    table.cell(1, 1).text = "+44 7400 123456"
    table.cell(1, 2).text = "café and 東京"
    table.cell(1, 2).add_paragraph("Next cell paragraph")
    document.add_paragraph("Meeting: 2026-10-03. Secret: ghp_Fictional7654.")
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(
        qn("r:id"),
        document.part.relate_to(
            "https://example.test/" + SENTINEL, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
        ),
    )
    link_run = OxmlElement("w:r")
    link_text = OxmlElement("w:t")
    link_text.text = "Fictional portal"
    link_run.append(link_text)
    hyperlink.append(link_run)
    document.add_paragraph()._p.append(hyperlink)

    def png_chunk(kind, data):
        return (
            struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))
        )

    image = (
        b"\x89PNG\r\n\x1a\n"
        + png_chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
        + png_chunk(b"IDAT", zlib.compress(b"\x00\x20\x80\x60"))
        + png_chunk(b"IEND", b"")
    )
    document.add_picture(io.BytesIO(image))
    document.sections[0].footer.paragraphs[0].text = "Fictional footer"
    hidden = OxmlElement("w:del")
    run = OxmlElement("w:r")
    text = OxmlElement("w:delText")
    text.text = SENTINEL
    run.append(text)
    hidden.append(run)
    person._p.append(hidden)
    buffer = io.BytesIO()
    document.save(buffer)
    output = io.BytesIO()
    with (
        ZipFile(io.BytesIO(buffer.getvalue())) as source,
        ZipFile(output, "w", ZIP_DEFLATED) as target,
    ):
        for name in source.namelist():
            data = source.read(name)
            if name == "customXml/item1.xml":
                data = f'<private xmlns="urn:fictional">{SENTINEL}</private>'.encode()
            target.writestr(name, data)
    return output.getvalue()


def simple_docx(*paragraphs):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def read_word(payload):
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(io.BytesIO(payload))
    chunks = []
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            chunks.append(block.text)
        elif isinstance(block, Table):
            chunks.extend(
                "\t".join("\n".join(p.text for p in cell.paragraphs) for cell in row.cells)
                for row in block.rows
            )
    return "\n".join(chunks), document
