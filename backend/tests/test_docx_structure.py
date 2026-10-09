import io
import json
from datetime import UTC, datetime
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from app.exports.docx import PARTS_ALLOWLIST, generate_word
from app.intake.imports import extract_import
from app.intake.structure import (
    InvalidLayout,
    allows_span,
    joined_text,
    leaves,
    realign_word,
    validate_layout,
)
from tests.docx_fixtures import SENTINEL, read_word, simple_docx, structured_docx

NOW = datetime(2026, 10, 3, 12, tzinfo=UTC)


def test_import_exact_order_spans_roles_restart_and_merged_cells():
    imported = extract_import("fictional.docx", structured_docx())
    source, layout = imported.source.text, imported.layout
    assert source.startswith("Header: nora@example.test\nFictional research\n")
    assert source.endswith("Fictional footer") and SENTINEL not in source
    assert joined_text(layout, source) == source
    blocks = list(leaves(layout))
    assert any(block["role"] == "title" for block in blocks)
    assert {block["level"] for block in blocks if block["role"] == "heading"} == {1, 9}
    assert {block["level"] for block in blocks if block.get("list_kind") == "bullet"} == {0, 2}
    numbers = [block for block in blocks if block.get("list_kind") == "number"]
    assert numbers[0]["list_id"] == numbers[1]["list_id"] != numbers[2]["list_id"]
    assert numbers[-1]["level"] == 8
    table = next(block for block in layout["blocks"] if block["kind"] == "table")
    assert [len(row) for row in table["rows"]] == [2, 3]
    assert "nora" not in json.dumps(layout) and "text" not in json.dumps(layout)


def test_fresh_word_readback_layout_neutral_properties_allowlist_and_no_hidden_sentinel():
    imported = extract_import("fictional.docx", structured_docx())
    payload = generate_word(imported.source.text, imported.layout, NOW)
    text, document = read_word(payload)
    assert text == imported.source.text
    assert document.paragraphs[1].style.name == "Title"
    assert any(p.style.name == "Heading 9" for p in document.paragraphs)
    assert any(p.style.name == "List Number 9" for p in document.paragraphs)
    nums = [p._p.pPr.numPr.numId.val for p in document.paragraphs if p.style.name == "List Number"]
    assert nums[0] == nums[1] != nums[2]
    numbering = document.part.numbering_part.element
    for number in set(nums):
        num = next(
            node
            for node in numbering.findall(qn("w:num"))
            if node.get(qn("w:numId")) == str(number)
        )
        assert (
            num.find("./" + qn("w:lvlOverride") + "/" + qn("w:startOverride")).get(qn("w:val"))
            == "1"
        )
    with ZipFile(io.BytesIO(payload)) as archive:
        assert set(archive.namelist()) == PARTS_ALLOWLIST
        for name in archive.namelist():
            data = archive.read(name)
            assert SENTINEL.encode() not in data
            assert b'TargetMode="External"' not in data
        app = etree.fromstring(archive.read("docProps/app.xml"))
        assert len(app) == 1 and app[0].text == "OpenAnonymi"
    core = document.core_properties
    for name in (
        "author",
        "title",
        "subject",
        "keywords",
        "comments",
        "category",
        "content_status",
        "identifier",
        "language",
        "version",
        "last_modified_by",
    ):
        assert getattr(core, name) == ""
    assert core.revision == 1 and core.last_printed is None
    assert core.created == NOW


@pytest.mark.parametrize(
    "source", ["😀 before\n\n東京 after\n", "\nFirst\r\nNext\tcell", "a\rb", "  source  "]
)
def test_plain_word_paragraphs_preserve_exact_codepoints_including_crlf(source):
    payload = generate_word(source, None, NOW)
    assert read_word(payload)[0] == source


@pytest.mark.parametrize(
    "paragraphs", [("", "\n😀 first", "end\n", ""), ("A\nB", "C\tD"), ("\ufeff😀 café", "東京")]
)
def test_final_strip_and_bom_spans_align(paragraphs):
    imported = extract_import("fictional.docx", simple_docx(*paragraphs))
    assert imported.layout is not None
    assert joined_text(imported.layout, imported.source.text) == imported.source.text


def test_complex_layout_falls_back_without_losing_source(monkeypatch):
    monkeypatch.setattr("app.intake.structure.MAX_BLOCKS", 3)
    imported = extract_import("fictional.docx", simple_docx("A", "B", "C", "D"))
    assert imported.layout is None and imported.source.text == "A\nB\nC\nD"
    assert "plain paragraphs" in " ".join(imported.notes)


def test_one_leaf_edits_shift_later_spans_and_other_edits_simplify():
    imported = extract_import("fictional.docx", simple_docx("😀 Nora", "Next block"))
    old = imported.source.text
    result = realign_word(imported.layout, old, old.replace("Nora", "Nora Caldwell"))
    assert joined_text(result, old.replace("Nora", "Nora Caldwell")) == old.replace(
        "Nora", "Nora Caldwell"
    )
    assert list(leaves(imported.layout))[1]["start"] == 7
    assert list(leaves(result))[1]["start"] == 16
    assert realign_word(imported.layout, old, old.replace("Nora", "Nora\tCaldwell")) is None
    assert realign_word(imported.layout, old, old.replace("\n", " ")) is None
    assert (
        realign_word(imported.layout, old, old.replace("😀", "X").replace("block", "note")) is None
    )


def test_crossing_and_control_boundaries_fail_but_unstructured_keeps_legacy_behavior():
    imported = extract_import("fictional.docx", simple_docx("Nora", "Caldwell", "A\nB", "C\tD"))
    source = imported.source.text
    assert allows_span(imported.layout, source, 0, 4)
    assert not allows_span(imported.layout, source, 0, 13)
    assert not allows_span(imported.layout, source, source.index("A"), source.index("A") + 3)
    assert not allows_span(imported.layout, source, source.index("C\t"), source.index("C\t") + 3)
    assert allows_span(None, source, 0, 13)
    broken = json.loads(json.dumps(imported.layout))
    broken["blocks"][0]["end"] = 1000
    with pytest.raises(InvalidLayout):
        validate_layout(broken, len(source), source)


@pytest.mark.parametrize(
    "extra", ["leading_cell", "trailing_cell", "leading_row", "trailing_row", "nested"]
)
def test_edge_table_blank_paragraphs_and_nested_table_flattening_keep_exact_map(extra):
    document = Document()
    table = document.add_table(rows=2 if "row" in extra else 1, cols=1)
    if extra == "nested":
        table.cell(0, 0).add_table(rows=1, cols=2).cell(0, 0).text = "Nested text"
    else:
        cell = table.cell(1 if extra == "leading_row" else 0, 0)
        if extra == "leading_cell":
            cell.add_paragraph("Fictional text")
        else:
            cell.text = "Fictional text"
        if extra == "trailing_cell":
            cell.add_paragraph("")
    stream = io.BytesIO()
    document.save(stream)
    imported = extract_import("fictional.docx", stream.getvalue())
    assert imported.layout is not None
    assert joined_text(imported.layout, imported.source.text) == imported.source.text
    assert (
        read_word(generate_word(imported.source.text, imported.layout, NOW))[0]
        == imported.source.text
    )


def test_layout_rejects_untyped_version_and_unused_fields():
    imported = extract_import("fictional.docx", simple_docx("Fictional text"))
    imported.layout["v"] = True
    with pytest.raises(InvalidLayout):
        validate_layout(imported.layout, len(imported.source.text))
    imported.layout["v"] = 1
    imported.layout["blocks"][0]["level"] = "untrusted text"
    with pytest.raises(InvalidLayout):
        validate_layout(imported.layout, len(imported.source.text))
