"""Small fictional documents, constructed in memory for real extraction checks."""

import io


def docx_sample():
    from docx import Document

    document = Document()
    document.core_properties.author = "Fictional metadata author"
    document.sections[0].header.paragraphs[0].text = "Fictional research note"
    document.add_paragraph("😀 Nora Caldwell; nora@example.com; café and 東京.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Case"
    table.cell(0, 1).text = "CASE-123456"
    document.add_paragraph("A final fictional paragraph.")
    document.sections[0].footer.paragraphs[0].text = "Fictional footer"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pdf_sample(*, encrypted=False, empty=False, pages=1, inflated=False):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    writer.add_metadata(
        {"/Author": "Fictional metadata author", "/Title": "Private fictional metadata"}
    )
    for _ in range(pages):
        page = writer.add_blank_page(width=612, height=792)
        if not empty:
            font = DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                }
            )
            page[NameObject("/Resources")] = DictionaryObject(
                {
                    NameObject("/Font"): DictionaryObject(
                        {NameObject("/F1"): writer._add_object(font)}
                    )
                }
            )
            stream = DecodedStreamObject()
            stream.set_data(
                (b" " * 1100000 if inflated else b"")
                + b"BT /F1 12 Tf 50 720 Td (Nora Caldwell; nora@example.com. A fictional PDF note.) Tj ET"
            )
            page[NameObject("/Contents")] = writer._add_object(
                stream.flate_encode() if inflated else stream
            )
    if encrypted:
        writer.encrypt("fictional-password")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()
